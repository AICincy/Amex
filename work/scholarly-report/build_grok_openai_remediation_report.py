from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).with_name("grok_openai_skill_remediation_technical_report.docx")

BLACK = "000000"
NAVY = "163A5F"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F5F6"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_table_widths(table, widths) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_paragraph(doc, text="", *, style=None, italic=False, bold_lead=None, align=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        r = p.add_run(text)
        r.italic = italic
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for index, value in enumerate(headers):
        cell = hdr.cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        fill = "FFFFFF" if row_index % 2 == 0 else PALE_BLUE
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_shading(cell, fill)
            set_cell_borders(cell)
            set_cell_margin(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            if index == 0 and len(str(value)) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
    if widths:
        set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.74)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.space_after = Pt(6)
    for name, size, before, after in (("Title", 18, 0, 12), ("Heading 1", 13.2, 15, 5), ("Heading 2", 11.4, 11, 4), ("Heading 3", 10.6, 8, 3)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Caption" in styles:
        styles["Caption"].font.name = "Times New Roman"
        styles["Caption"].font.size = Pt(8.7)
        styles["Caption"].font.italic = True
        styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("PREPRINT TECHNICAL REPORT   |   16 SEPTEMBER 2026")
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = footer.add_run("Page ")
    f_run.font.name = "Times New Roman"
    f_run.font.size = Pt(8)
    add_field(footer, "PAGE")


def main() -> None:
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Remediating Grok Origin Skills for OpenAI Compatible Codex Deployment")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A provenance preserving technical account of package translation controls and verification")
    r.italic = True
    r.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Technical remediation report  |  Evidence date 16 September 2026")
    r.font.size = Pt(9.4)
    p.paragraph_format.space_after = Pt(15)

    abs_label = doc.add_paragraph()
    abs_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = abs_label.add_run("Abstract")
    r.bold = True
    r.font.size = Pt(11.5)
    abstract = (
        "This report reconstructs and evaluates the remediation of a Grok-origin skill workspace for use as local, OpenAI-compatible Codex skill packages. "
        "The remediation was not a mechanical rename. It established a portable package contract, constrained host binding to observed Codex capabilities, separated runtime instructions from mutable world knowledge, added authorization and credential boundaries, and introduced deterministic evidence collection. "
        "The repository lineage records package phases for claim auditing, Reddit governance, API and drafting operations, and prompt architecture. A recorded 27-case suite passed at the repository evidence point. A fresh re-execution in the present Codex environment produced 26 passes and one expected-removal failure because a local `resend-api` package remains installed despite the repository removal. "
        "Accordingly, the defensible conclusion is conditional: the repository remediation is documented and its retained controls are locally testable, but current environment conformity is incomplete until that stale local package is removed and the suite is re-run. This report does not assert ChatGPT activation, platform certification, or universal runtime safety."
    )
    p = add_paragraph(doc, abstract, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Keywords  ")
    r.bold = True
    p.add_run("Codex skills, OpenAI metadata, host portability, authorization gates, provenance, adversarial evaluation, static validation")
    p.paragraph_format.space_after = Pt(14)

    add_heading(doc, "1  Scope and evidentiary posture")
    add_paragraph(doc, "The object of analysis is a local repository whose `.grok/workspace.yaml` identifies the source host as `grok` [R1]. The target is not an assertion that a third-party system is OpenAI-certified. Rather, it is a narrow implementation claim: the packages were reshaped to use a Codex-oriented skill layout and to obey an explicit control contract compatible with the current local host. The phrase OpenAI compatible in this report therefore means package and behavioral compatibility with the local Codex conventions represented in the repository, not a warranty, attestation, or approval by OpenAI.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_paragraph(doc, "Claims are stratified by evidence type. Repository history establishes what was committed; source files establish what instructions and helper code specify; a deterministic run establishes only the checks executed in that run. Neither a YAML manifest, a static gate, nor a local dry run establishes a skill's installation on another product, activation in ChatGPT, completeness under all prompts, or safety outside the tested threat model [R2, R3].")
    add_table(doc, ["Evidence class", "What it supports", "What it does not support"], [
        ["Git lineage", "A remediation change was committed with a timestamp and file set", "Current installation or live behavior"],
        ["Package source", "Declared controls, metadata, references, and helper logic", "Host loading, credentials, external account state"],
        ["Static package gate", "Deterministic structural conformance of its target", "Activation, production effect, adversarial generalization"],
        ["Local behavioral case", "Observed helper response for one credential-free input", "Provider response, billing outcome, broad runtime reliability"],
        ["Host manifest", "A request for Codex product binding", "That the host installed, invoked, or activated the package"],
    ], [1.16, 2.65, 2.68])

    add_heading(doc, "2  System model and remediation objective")
    add_paragraph(doc, "The original workspace marker provided provenance, not a portable runtime contract. The remediation therefore treated the skill family as a control system with four separable planes: package identity, active-turn coordination, external-action control, and evidence. This decomposition prevents a platform-specific archive from becoming an unreviewable source of authority after migration.")
    add_table(doc, ["Stage", "Controlled transformation", "Claim retained at the stage boundary"], [
        ["1  Provenance", "Record source-host marker and commit lineage.", "The source package and migration history are identifiable."],
        ["2  Package", "Use canonical skill names, `SKILL.md`, version records, and Codex metadata.", "The package describes an intended local Codex binding."],
        ["3  Control", "Constrain exposed tools, live sources, current authority, and data retention.", "The package does not carry historical host assumptions as authority."],
        ["4  Execution", "Introduce dry runs, explicit execute flags, origin controls, and human gates.", "A tested helper either remains local or blocks before a live action."],
        ["5  Evidence", "Record static and local behavioral outcomes with stated limitations.", "A narrow result can be reproduced without claiming host activation."],
    ], [0.88, 3.36, 2.24], 8.0)
    add_paragraph(doc, "The objective was to retain host-neutral value: task decomposition, source provenance, safe failure reporting, human gates, and deterministic tooling. It was equally important to discard or quarantine host-specific authority: Grok paths, presumed connectors, hidden inter-skill channels, stored account state, implicit permission, and dated fact tables. In formal terms, the migration maps a source package S to a target package T while reducing the authority set A(S) to the least set A(T) that the visible Codex turn can justify.")
    p = add_paragraph(doc, "A(T) = { instructions, local files, exposed tools, current authorization, and current source evidence }", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=5, after=8)
    add_paragraph(doc, "All other candidate capabilities are treated as absent until discovered. This is a fail-closed portability rule. It prevents a static archive from laundering historical access or an old integration name into a current operational claim.")

    add_heading(doc, "3  Remediation method")
    add_heading(doc, "3.1  Package normalization and identity control", 2)
    add_paragraph(doc, "Each remediated package was expressed as a named skill directory with a `SKILL.md` contract, a `VERSION` record, and Codex-facing `agents/openai.yaml` metadata. The metadata exposes a display contract and target product policy but repeatedly disclaims activation. Canonical directory names are specified so that transient export folders cannot be mistaken for persistent package identity [R2, R4]. This distinction is material because installation, archive presence, and runtime invocation are different observables.")
    add_heading(doc, "3.2  Governor and active turn coordination", 2)
    add_paragraph(doc, "The remediation introduced a governing layer, `aai-cognitive-interface`, and made relevant domain packages subordinate to it. The resulting model prevents domain modules from self-governing, inventing background processes, or treating an internal sibling as a privileged channel. In the claim-source package, for example, siblings are only candidate evidence; the live agent coordinates only tools exposed in the current turn [R4]. The Reddit operations package similarly declares that it cannot execute a moderator action and must stop at the human gate for live mutations [R5].")
    add_heading(doc, "3.3  Runtime only knowledge boundary", 2)
    add_paragraph(doc, "The packages were converted from possible repositories of remembered operational facts into runtime overlays. The AAI contract prohibits using bundled statutes, addresses, watch lists, dated sources, account state, or matter facts as current knowledge. Such content must be re-fetched from a current primary source or marked unresolved [R3]. The control matters for scientific reproducibility as well as safety: execution policy is versioned, while empirical input remains contemporaneous and separately attributable.")
    add_heading(doc, "3.4  Host binding and status taxonomy", 2)
    add_paragraph(doc, "Host bindings identify Codex as the intended product and enumerate a local persist path, yet expressly reject the inference that metadata proves installation or activation. The package identity contract permits a `STATIC PASS` conclusion only when a named gate succeeds; it withholds saved, installed, smoke, runtime-verified, and adversarial labels absent stronger evidence [R2, R3]. This removes a common migration defect in which a manifest is rhetorically upgraded into proof of behavior.")

    add_heading(doc, "4  External action remediation")
    add_paragraph(doc, "The most security-relevant changes occurred where the skills could access accounts, send requests, charge an account, disclose a credential, or affect third parties. The chosen design is not merely advisory. Helper behavior was rewritten so the default path is inspectable and credential-free, and a live operation requires an additional explicit gate. The following matrix maps implementation changes to the corresponding failure mode.")
    add_table(doc, ["Module", "Remediation", "Operational consequence", "Evidence"], [
        ["Dify Agent API", "Added `--dry-run` and requires `--execute` before a live request; dry run returns a request plan without loading credentials.", "Prevents accidental network calls during inspection and makes live invocation mechanically distinct.", "Commit 8aa09d1 diff; current dynamic cases."],
        ["EnformionGO API", "Removed live HTTP execution and credential loading; helper now blocks live calls pending a scope-bound authorization receipt.", "Prevents a potentially billed sensitive-record search from being initiated by package code.", "Commit 8aa09d1 diff; current live-gate case."],
        ["Speko API", "Requires `--execute`; validates relative paths; pins origin to HTTPS `api.speko.dev`; rejects redirects and unsafe methods.", "Reduces credential forwarding, origin-confusion, and accidental action risks.", "Commit 8aa09d1 diff; current origin-gate case."],
        ["Resend API", "Repository package removed by approved remediation rather than retained as a latent send route.", "The target state contains no repository-managed Resend skill. Current local residue remains an exception.", "Commit 8aa09d1; fresh suite failure."],
        ["Reddit Owner Ops", "Converts live moderation actions into named human gates and separates drafting from publication.", "A drafted artifact is not represented as a completed Reddit write.", "Package contract [R5]."],
    ], [1.18, 2.15, 2.10, 1.05], 7.6)
    add_paragraph(doc, "The Enformion change deserves special notice. The earlier helper performed an HTTP request after loading access-profile material. The remediated helper retains a dry-run plan but terminates a non-dry path with a block message. This is defense in depth: the natural-language instruction says live execution is disabled, and the helper implements that restriction. Likewise, Speko's URL validation precedes credential loading, so an attacker-controlled destination cannot obtain the authorization header through a malformed path or redirect.")

    add_heading(doc, "5  Scientific and evidentiary controls")
    add_heading(doc, "5.1  Claim source audit discipline", 2)
    add_paragraph(doc, "The claim-source auditor formalizes a non-blended taxonomy: verified, conflicting, not found in searched sources, and manual review needed. Each claim is assigned exactly one status, and citation presence alone is declared insufficient for entailment [R4]. This turns fact checking from an informal confidence judgment into an auditable row-level method. It also avoids a recurrent scientific-reporting error: treating absence of a source within a specified corpus as proof of falsity.")
    add_heading(doc, "5.2  Source hierarchy and provenance", 2)
    add_paragraph(doc, "The skill family treats mutable reference data as untrusted cache. A live source, an authorized matter file, or present-turn retrieval is required for current factual assertions. This prevents source-host archives from carrying stale legal, operational, or personal-data assertions into a different deployment. Package documentation is thus a methods specification, not a facts database.")
    add_heading(doc, "5.3  Failure visibility and no simulated completion", 2)
    add_paragraph(doc, "The governing contract requires failures, empty returns, missing capabilities, and blocked operations to remain visible. It also distinguishes draft, static, saved, installed, smoke, runtime-verified, and adversarial states. These distinctions were carried into the report because they limit the epistemic scope of the result. A well-designed skill should fail legibly rather than narrate success after a tool or credential boundary prevented execution.")

    add_heading(doc, "6  Implementation lineage")
    add_paragraph(doc, "Repository history supplies the reproducible change sequence. The listed commits are not used as evidence of external deployment; they show the local evolution of the remediation. The breadth of the file sets demonstrates that the work included package interfaces, references, schemas, fixtures, helper programs, and tests rather than only prompt text.")
    add_table(doc, ["Commit", "Recorded change", "Primary effect"], [
        ["874a4bd", "Add remediated claim source auditor skill", "Adds an evidence taxonomy, package gate, claim schema, and source-audit controls."],
        ["f7b9b45", "Add governed Reddit skill phase", "Adds controlled drafting, AutoMod auditing, public-copy scanning, and explicit human gates."],
        ["28ac242", "Add governed API and drafting skill phase", "Adds Dify, Enformion, Speko, CourtListener, research, complaint, record, and mediation packages."],
        ["8aa09d1", "Harden skills and add adversarial suite", "Adds dry-run and execute controls, eliminates repository Resend route, and records 27 local checks."],
        ["220896f", "Add scholarly critical thinking skill prompts", "Adds 20 adversarial prompt artifacts emphasizing epistemic self-audit."],
    ], [1.10, 2.45, 2.93])
    add_paragraph(doc, "The hardened API diff provides concrete examples of the migration's technical direction. Dify read routes changed from immediate credential-backed network requests to a parser with mutually recognizable dry-run and execute states. Enformion's network and environment-loading code was deleted from the routine execution path. Speko added a relative-path parser, permitted-method allowlist, canonical origin, and a redirect-denying opener. These changes reduce the privileges attached to a prompt-produced command.")

    add_heading(doc, "7  Verification protocol and observed results")
    add_paragraph(doc, "Verification has two timepoints. First, the repository's saved report records a 27-case pass on 16 September 2026, including 20 prompt-to-control-marker checks and seven local behavioral cases [R6]. Second, this report re-executed the same suite in the current Codex environment on 16 September 2026 at 21:02 UTC using a temporary output path. The re-run is important because remediation state is not equivalent to current installation state.")
    add_table(doc, ["Evaluation", "Recorded result", "Meaning", "Boundary"], [
        ["Repository saved suite", "27 of 27 PASS", "Historical evidence that retained controls and seven no-network helpers passed at that evidence point.", "Static controls and credential-free local behavior only."],
        ["Fresh Codex re-run", "26 PASS, 1 FAIL", "All retained prompt cases and all seven dynamic cases passed; `resend-api` presence violated the approved-removal expectation.", "Environment conformity is incomplete."],
        ["AAI package gate", "PASS, 0 violations", "The staged AAI package satisfied its own static package validator with user Python 3.14.6.", "Static structure only."],
        ["Subordinate staged gates", "BLOCKED in staging", "Subordinate validators searched for the canonical installed AAI gate and did not find it in the staging path.", "Not a package pass or fail; installed-path dependency unresolved in staging."],
    ], [1.31, 1.23, 2.46, 1.48], 7.7)
    add_heading(doc, "7.1  Fresh discrepancy and its interpretation", 2)
    add_paragraph(doc, "The sole fresh-suite failure is not a test flake. The suite expects `C:\\Users\\jared\\.codex\\skills\\resend-api\\SKILL.md` to be absent because the repository remediation removed the staged Resend package. The file is presently installed in the local skills root. Therefore the current environment contains a stale local package that the repository target state does not authorize. The correct result is a recorded configuration discrepancy, not a claim that the broader remediation failed or that the local package has necessarily performed an external action.")
    add_paragraph(doc, "The exact corrective sequence is narrow: remove or archive the local `resend-api` skill only under explicit authorization, then re-run `scripts/run_adversarial_skill_suite.py` and confirm 27 of 27 pass. This report does not perform that deletion because its authorized objective is explanation and document production, not mutation of the installed skill root.")

    add_heading(doc, "8  Residual risk and non claims")
    add_paragraph(doc, "A high-quality remediation report must make its negative space explicit. The controls described here do not prove that OpenAI has certified, reviewed, or endorsed these skills. They do not prove that ChatGPT has loaded any local Codex directory. They do not guarantee correct model interpretation for all prompts, assure security of downstream providers, establish account permissions, or validate unaudited dependencies. They do not substitute for a current source review before a consequential provider call.")
    add_paragraph(doc, "The remaining risk is managed partly by design and partly by process. Design restrictions include constrained helper behavior, origin validation, secret non-persistence, status taxonomy, and no-live-write gates. Process restrictions include a current-turn authorization requirement, human gates for irreversible actions, source re-fetch before consequential changes, and evidence collection after material revisions. This division is intentional: no static skill file can fully authorize or verify live activity outside the host context.")
    add_table(doc, ["Risk class", "Implemented mitigation", "Residual condition"], [
        ["Host-assumption drift", "Use only currently exposed tools and label host metadata as non-activation proof.", "Host inventory can still change between turns."],
        ["Credential disclosure", "No keys in packages; dry-run before credential loading; canonical-origin validation.", "External secret stores and providers remain outside this report."],
        ["Unauthorized external effect", "Human gates, `--execute`, and Enformion live disablement.", "An authorized user may later choose to execute an action."],
        ["Stale facts", "Runtime-only rule and primary-source re-fetch requirement.", "A later agent must actually perform the required retrieval."],
        ["Configuration drift", "Deterministic suite and expected-removal test.", "Fresh run presently identifies a stale local Resend package."],
    ], [1.42, 2.62, 2.43])

    add_heading(doc, "9  Conclusion")
    add_paragraph(doc, "The remediation converted a Grok-origin workspace into a Codex-oriented skill family by making authority explicit, reducing implicit host assumptions, separating current facts from reusable methods, and encoding irreversible-action boundaries in both instructions and helper behavior. The strongest technical contribution is not the presence of OpenAI metadata alone. It is the closed chain from canonical package identity, through active-turn and source controls, to deterministic local evidence and unambiguous status labels.")
    add_paragraph(doc, "At the repository level, the recorded remediation is internally coherent and its saved local suite passed. At the current environment level, one stale local `resend-api` directory prevents a full-conformity conclusion. The report therefore reaches a qualified scientific conclusion: the remediation design and retained testable controls are evidenced; the destination installation requires one configuration reconciliation before the current 27-case suite can be claimed as passing in this Codex environment.")

    add_heading(doc, "Appendix A  Remediated skill inventory")
    inventory = [
        ("aai-cognitive-interface", "Governing state custody and evidence controls"),
        ("authority-currency-auditor", "Primary-source authority and currency analysis"),
        ("claim-source-auditor", "Atomic claim and evidence reconciliation"),
        ("courtlistener-api", "Controlled legal-research API method"),
        ("dify-agent-api", "Dify service API with dry-run and execute gates"),
        ("enformion-go-api", "Credential-free query planning with live disablement"),
        ("exa-firecrawl", "Current web-retrieval method boundary"),
        ("forensic-evidentiary-drafting", "Evidentiary drafting method"),
        ("personal-context", "Prior-artifact retrieval with scope controls"),
        ("practitioner-narrative-writer", "Professional narrative drafting control"),
        ("prompt-architecture-engineering", "Prompt design without invented tool claims"),
        ("record-series-builder", "Record series, indexes, and state control"),
        ("reddit-automod-yaml", "Static AutoMod analysis and public-copy controls"),
        ("reddit-owner-ops", "Drafting and owner-operation human gates"),
        ("register-mediation", "One-retry mediation workflow"),
        ("regulatory-complaint-drafting", "Evidence-bound regulatory drafting"),
        ("research-execution-briefs", "Scoped research execution and source hierarchy"),
        ("resend-api", "Approved repository removal; stale local package detected"),
        ("speko-mcp", "Origin-constrained API helper with dry-run"),
        ("subreddit-rule-packet", "Rule packet generation and validation"),
    ]
    add_table(doc, ["Skill", "Remediated control role"], inventory, [2.25, 4.22], 8.4)

    add_heading(doc, "Appendix B  Reproduction procedure")
    add_paragraph(doc, "The following procedure reproduces the no-network evidence collection used by the suite. It is intentionally distinct from a provider test and should be executed only in a controlled local environment. The present report used a temporary JSON output path so that the repository audit record was not overwritten.")
    code = [
        "python scripts/run_adversarial_skill_suite.py --output <temporary-report>.json",
        "python work/aai-cognitive-interface/scripts/aai_runtime_gate.py package work/aai-cognitive-interface",
        "Inspect <temporary-report>.json for prompt mappings, each expected exit code, and any removal-state discrepancy.",
        "Do not infer installation, activation, external delivery, or provider-side success from these outputs.",
    ]
    for line in code:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("•  " + line)
        r.font.name = "Courier New" if line.startswith("python") else "Times New Roman"
        r.font.size = Pt(9.2)

    add_heading(doc, "References")
    refs = [
        "[R1] `.grok/workspace.yaml`, repository marker identifying `host: grok` and the workspace identity. Inspected 16 September 2026.",
        "[R2] `work/aai-cognitive-interface/SKILL.md` and `references/package-identity.md`, control contract, canonical identity, and status restrictions. Inspected 16 September 2026.",
        "[R3] `work/aai-cognitive-interface/references/host-binding.md`, `runtime-only.md`, and `trusted-controller.md`, host and runtime claim boundaries. Inspected 16 September 2026.",
        "[R4] `.aai-stage-claim-source-auditor/claim-source-auditor/SKILL.md` and `agents/openai.yaml`, source-audit and active-turn coordination contract. Inspected 16 September 2026.",
        "[R5] `.aai-stage-reddit-governed-skills/reddit-owner-ops/SKILL.md` and `agents/openai.yaml`, live-mutation human gates and Codex binding. Inspected 16 September 2026.",
        "[R6] `audits/adversarial-skill-suite-2026-09-16.json` and `scripts/run_adversarial_skill_suite.py`, recorded 27-case suite and its deterministic methodology. Inspected and re-run 16 September 2026.",
        "[R7] Git commits `874a4bd`, `f7b9b45`, `28ac242`, `8aa09d1`, and `220896f`, remediation lineage and change sets. Inspected 16 September 2026.",
        "[R8] Current temporary run of `scripts/run_adversarial_skill_suite.py` at 2026-09-16T21:02:02Z. Result: 26 passed, 1 failed because local `resend-api` remains present.",
    ]
    for ref in refs:
        p = add_paragraph(doc, ref, after=4)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        for run in p.runs:
            run.font.size = Pt(8.8)

    doc.core_properties.title = "Remediating Grok Origin Skills for OpenAI Compatible Codex Deployment"
    doc.core_properties.subject = "Technical remediation report"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Codex, OpenAI, Grok, skills, remediation, verification"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
