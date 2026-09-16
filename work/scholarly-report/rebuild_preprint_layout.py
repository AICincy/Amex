from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("grok_openai_skill_remediation_technical_report.docx")
NAVY = "1F4E79"
PALE = "F2F6FA"
BORDER = "C8D0D8"


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = pr.find(qn("w:shd"))
    if el is None:
        el = OxmlElement("w:shd")
        pr.append(el)
    el.set(qn("w:fill"), fill)


def border(cell):
    pr = cell._tc.get_or_add_tcPr()
    edges = pr.first_child_found_in("w:tcBorders")
    if edges is None:
        edges = OxmlElement("w:tcBorders")
        pr.append(edges)
    for side in ("top", "left", "bottom", "right"):
        edge = edges.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            edges.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), BORDER)


def margin(cell):
    pr = cell._tc.get_or_add_tcPr()
    mar = pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        pr.append(mar)
    for side, val in (("top", 80), ("bottom", 80), ("start", 105), ("end", 105)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def width(cell, inches):
    cell.width = Inches(inches)
    pr = cell._tc.get_or_add_tcPr()
    tcw = pr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        pr.append(tcw)
    tcw.set(qn("w:w"), str(int(inches * 1440)))
    tcw.set(qn("w:type"), "dxa")


def set_repeat(row):
    pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    pr.append(el)


def keep(p):
    pr = p._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    pr.append(el)


def run_font(run, size=None, bold=None, italic=None, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def p(doc, text="", style=None, justify=True, size=None, italic=False, before=None, after=None):
    para = doc.add_paragraph(style=style)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    r = para.add_run(text)
    run_font(r, size=size, italic=italic)
    if before is not None:
        para.paragraph_format.space_before = Pt(before)
    if after is not None:
        para.paragraph_format.space_after = Pt(after)
    return para


def heading(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    para.add_run(text)
    keep(para)
    return para


def caption(doc, text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after = Pt(3)
    r = para.add_run(text)
    run_font(r, size=8.7, italic=True)
    keep(para)


def table(doc, headers, rows, widths, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    h = t.rows[0]
    set_repeat(h)
    for i, item in enumerate(headers):
        c = h.cells[i]
        width(c, widths[i]); shade(c, NAVY); border(c); margin(c)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        r = para.add_run(item); run_font(r, size=font_size, bold=True)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for row_number, items in enumerate(rows):
        cells = t.add_row().cells
        for i, item in enumerate(items):
            c = cells[i]
            width(c, widths[i]); shade(c, "FFFFFF" if row_number % 2 == 0 else PALE); border(c); margin(c)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = c.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(str(item)) <= 18 else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            r = para.add_run(str(item)); run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def page(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.84); sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(0.92); sec.right_margin = Inches(0.92)
    sec.header_distance = Inches(0.30); sec.footer_distance = Inches(0.35)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.12; normal.paragraph_format.space_after = Pt(6)
    for style_name, size, b, a in (("Title", 18, 0, 12), ("Heading 1", 13, 16, 5), ("Heading 2", 11.2, 10, 3)):
        s = doc.styles[style_name]
        s.font.name = "Times New Roman"; s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(b); s.paragraph_format.space_after = Pt(a); s.paragraph_format.keep_with_next = True
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = foot.add_run("Preprint technical report   |   16 September 2026   |   ")
    run_font(r, size=8)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    foot._p.append(fld)


def main():
    doc = Document(); setup(doc)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Remediating Grok Origin Skills for OpenAI Compatible Codex Deployment")
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("A provenance preserving technical account of package translation controls and verification"); run_font(r, size=11, italic=True)
    by = doc.add_paragraph(); by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = by.add_run("Technical remediation report   |   Evidence date 16 September 2026"); run_font(r, size=9.2)
    by.paragraph_format.space_after = Pt(16)
    label = doc.add_paragraph(); label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = label.add_run("Abstract"); run_font(r, size=11, bold=True)
    abstract = ("This report reconstructs the migration of a Grok-origin skill workspace into local, OpenAI-compatible Codex skill packages. The work established canonical package identity, explicit Codex metadata, active-turn coordination, runtime-only knowledge boundaries, authorization gates, credential-safe helper behavior, and deterministic evidence collection. It was not a mechanical rename. The implementation deliberately preserved reusable methods while discarding host-specific assumptions, hidden capability claims, and archived operational facts. The repository records a 27-case local pass; a fresh Codex re-run produced 26 passes and one expected-removal failure because a local `resend-api` package is still installed even though the repository target removes it. The defensible conclusion is therefore conditional: the remediation architecture and retained testable controls are evidenced, but current environment conformity remains incomplete until that stale package is reconciled. This document does not claim ChatGPT activation, OpenAI certification, or universal runtime safety.")
    ap = p(doc, abstract, justify=True, after=10); ap.paragraph_format.left_indent = Inches(0.20); ap.paragraph_format.right_indent = Inches(0.20)
    kw = doc.add_paragraph(); kw.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a = kw.add_run("Keywords  "); run_font(a, size=9.2, bold=True)
    b = kw.add_run("Codex skills, OpenAI metadata, host portability, authorization gates, provenance, adversarial evaluation"); run_font(b, size=9.2)

    page(doc)
    heading(doc, "1  Scope and evidence boundary")
    p(doc, "The repository marker `.grok/workspace.yaml` identifies the source host as `grok` [R1]. The target condition is not vendor certification. It is a narrower and testable proposition: the packages were reshaped for local Codex conventions and are constrained by controls that are compatible with the visible local host. Consequently, OpenAI compatible in this report means an evidence-bounded implementation posture, not approval, endorsement, or a guarantee by OpenAI.")
    p(doc, "The report uses three evidence classes. Git history establishes that a change was committed. Package source establishes the declared instructions and helper behavior. A deterministic execution establishes only the local checks performed at that time. Neither a manifest nor a static pass proves installation on another product, activation in ChatGPT, provider-side effect, or generalized model behavior [R2, R3].")
    caption(doc, "Table 1. Evidence interpretation rules")
    table(doc, ["Evidence", "Defensible conclusion", "Excluded conclusion"], [
        ["Git lineage", "A specific remediation change was committed.", "Current installation or live behavior."],
        ["Package source", "The package declares a bounded control contract.", "Host invocation or account authority."],
        ["Static gate", "The targeted structure passed deterministic checks.", "Activation, production effect, or broad reliability."],
        ["Local behavior", "One helper response was observed for a no-network case.", "Provider response, billing outcome, or external delivery."],
    ], [1.15, 2.70, 2.65], 8.5)

    page(doc)
    heading(doc, "2  System model and remediation objective")
    p(doc, "The migration treats a skill family as a control system rather than a collection of prompts. Four planes must remain separable: package identity, active-turn coordination, external-action control, and evidence. The source workspace contributes provenance, but it cannot carry forward authority merely because it existed on a former host.")
    caption(doc, "Table 2. Provenance preserving remediation pipeline")
    table(doc, ["Stage", "Controlled transformation", "Resulting claim"], [
        ["1  Provenance", "Record source-host marker and commit lineage.", "The source and migration history are identifiable."],
        ["2  Package", "Use canonical names, `SKILL.md`, version records, and Codex metadata.", "The package describes an intended local Codex binding."],
        ["3  Control", "Limit tools, sources, authority, and retained data to the active turn.", "Historical host assumptions do not become operational authority."],
        ["4  Execution", "Add dry runs, explicit execute gates, origin checks, and human gates.", "A helper either remains local or blocks before a live action."],
        ["5  Evidence", "Record static and local behavioral outcomes with their limits.", "A narrow result can be reproduced without an activation claim."],
    ], [1.10, 3.20, 2.20], 8.2)
    p(doc, "Formally, the target authority set is reduced to: instructions, local files, currently exposed tools, current authorization, and current source evidence. All other capabilities are treated as unavailable until independently discovered. This fail-closed construction prevents a source archive from laundering historical connectors, accounts, permissions, or facts into a new runtime.")

    page(doc)
    heading(doc, "3  Remediation method")
    heading(doc, "3.1  Package normalization and identity control", 2)
    p(doc, "The packages were expressed as named skill directories with a `SKILL.md` contract, a `VERSION` record, and Codex-facing `agents/openai.yaml` metadata. The metadata identifies product targeting and default prompts but explicitly disclaims activation. Canonical names prevent transport folders from being treated as a persistent package identity [R2, R4].")
    heading(doc, "3.2  Governor and active-turn coordination", 2)
    p(doc, "Relevant domain packages became subordinate to `aai-cognitive-interface`. That governor owns evidence boundaries, continuation, recovery, and status labels. A domain skill may provide a method, but it may not self-govern, invent an unseen sibling channel, turn a draft into an external act, or claim a tool that the current turn does not expose. The claim-source package treats sibling output only as candidate evidence [R4]. The Reddit owner-operations package treats every live moderation change as a human gate [R5].")
    heading(doc, "3.3  Runtime-only knowledge boundary", 2)
    p(doc, "The packages are runtime overlays, not retained factual databases. Bundled statutes, addresses, account data, watch lists, dated guidance, and matter facts are untrusted cache. A consequential factual claim must be re-fetched from a primary source or marked unresolved [R3]. This control prevents portability from becoming stale-data propagation.")
    heading(doc, "3.4  Status taxonomy and host binding", 2)
    p(doc, "Local host metadata is treated as a request for binding, not evidence that binding occurred. The package gate may support a static result after its named validator passes. It does not establish saved, installed, activated, smoke-tested, runtime-verified, or adversarial status. That separation removes the common defect of promoting a manifest into a behavior claim [R2, R3].")

    page(doc)
    heading(doc, "4  Controlled external actions")
    p(doc, "The highest-risk remediation occurred where a helper could invoke a provider, consume paid resources, disclose credentials, or act on a third party. The new pattern is behavioral, not rhetorical: dry-run first, explicit live gate second, current authorization always, and provider-specific restrictions where the risk warrants them.")
    caption(doc, "Table 3. Executable safety changes")
    table(doc, ["Module", "Implemented control", "Security and operations effect"], [
        ["Dify Agent API", "Added `--dry-run`; a live request requires `--execute`.", "Inspection does not load credentials or contact Dify; live action is mechanically distinct."],
        ["EnformionGO API", "Removed live HTTP execution and credential loading from helper behavior.", "A potentially billed sensitive-record query blocks pending a scope-bound authorization receipt."],
        ["Speko API", "Requires `--execute`; validates a relative path; pins HTTPS origin; rejects redirects.", "Reduces credential forwarding, origin confusion, and accidental action routes."],
        ["Resend API", "Removed from the repository target state.", "No repository-managed email-send route remains; a stale local installation is separately detected."],
        ["Reddit Owner Ops", "Explicit human gate for every live Reddit mutation.", "A draft or local YAML file cannot be represented as a completed platform write."],
    ], [1.25, 2.35, 2.90], 8.2)
    p(doc, "Two changes demonstrate the difference between policy language and executable control. Enformion does not merely instruct the agent to avoid a live search: the revised helper terminates a live path with a block response. Speko validates the relative path and canonical origin before any credential is read, then denies redirects. Those properties reduce the privileges attached to prompt-produced arguments.")

    page(doc)
    heading(doc, "5  Evidence and implementation lineage")
    heading(doc, "5.1  Claim source audit discipline", 2)
    p(doc, "The claim-source auditor assigns each atomic claim exactly one status: verified, conflicting, not found in searched sources, or manual review needed. Citation presence is not treated as entailment. A row-level taxonomy prevents blended confidence statements and preserves the difference between a false claim, a conflicting source, and an insufficiently searched corpus [R4].")
    heading(doc, "5.2  Failure visibility", 2)
    p(doc, "The governing contract retains errors, empty results, blocked actions, and unavailable tools as visible state. It prohibits simulated completion and distinguishes draft, static, saved, installed, smoke, runtime-verified, and adversarial labels. This is an evidentiary design decision: the reported state must never exceed the operation evidence supporting it.")
    caption(doc, "Table 4. Repository remediation lineage")
    table(doc, ["Commit", "Recorded change", "Material contribution"], [
        ["874a4bd", "Remediated claim-source auditor", "Adds claim taxonomy, source reconciliation, package gate, schema, and fixtures."],
        ["f7b9b45", "Governed Reddit skill phase", "Adds AutoMod auditing, public-copy controls, and live-action human gates."],
        ["28ac242", "Governed API and drafting phase", "Adds Dify, Enformion, Speko, CourtListener, research, complaint, record, and mediation modules."],
        ["8aa09d1", "Harden skills and add adversarial suite", "Adds dry-run and execute controls, removes repository Resend, and records local evidence."],
        ["220896f", "Scholarly critical-thinking prompts", "Adds 20 adversarial prompt artifacts emphasizing epistemic self-audit."],
    ], [1.05, 2.25, 3.20], 8.3)

    page(doc)
    heading(doc, "6  Verification and current discrepancy")
    p(doc, "The report retains both historical and present-time results because they answer different questions. The repository audit file records 27 local cases passing on 16 September 2026. The current report re-ran the same suite in the local Codex environment at 21:02 UTC and observed 26 passes and one failure. The difference is configuration drift, not a reason to overwrite the historical record [R6, R8].")
    caption(doc, "Table 5. Verification results and limits")
    table(doc, ["Evaluation", "Observed result", "Interpretation"], [
        ["Saved repository suite", "27 of 27 PASS", "Historical evidence that controls and no-network local helpers passed at the recorded time."],
        ["Fresh Codex suite", "26 PASS and 1 FAIL", "All retained control and dynamic cases passed; expected Resend removal did not match current local installation."],
        ["AAI package gate", "PASS with 0 violations", "The staged AAI package passed its deterministic structural validation with user Python 3.14.6."],
        ["Staged subordinate gates", "BLOCKED in staging", "Their shared canonical AAI gate cannot be resolved from the staging path; this does not prove a package pass or failure."],
    ], [1.45, 1.55, 3.50], 8.3)
    heading(doc, "6.1  Resend reconciliation", 2)
    p(doc, "The fresh failure is explicit. The suite expects `C:\\Users\\jared\\.codex\\skills\\resend-api\\SKILL.md` to be absent because commit `8aa09d1` removed the repository-managed Resend route. It is present in the local skill root. The correct conclusion is that the local environment contains a stale package not represented by the remediated repository target state. It is not evidence that an external email was sent, nor that the other controls failed.")
    p(doc, "The corrective action is intentionally narrow: under explicit authorization, remove or archive that local package and re-run the suite. This report does not perform the deletion because its authorized scope is documentation, not mutation of the installed skill root.")

    page(doc)
    heading(doc, "7  Residual risk and conclusion")
    p(doc, "The remediation provides an auditable engineering posture, not an absolute security claim. It does not prove that OpenAI has certified or reviewed the packages. It does not prove ChatGPT activation from local Codex files. It does not validate downstream provider permissions, account state, provider behavior, untested dependencies, or all model responses. Each consequential provider action remains subject to current authorization and current-source review.")
    caption(doc, "Table 6. Remaining risk after remediation")
    table(doc, ["Risk", "Mitigation", "Residual condition"], [
        ["Host drift", "Use current exposed tools and deny activation inference from metadata.", "Host capabilities can change after a package is written."],
        ["Credential disclosure", "No keys in package; dry-run first; canonical-origin controls.", "Secret stores and external providers remain outside the package."],
        ["Unauthorized effect", "Human gates, explicit execute flags, and Enformion live disablement.", "A later authorized operation may still occur."],
        ["Stale facts", "Runtime-only rule plus required current primary-source retrieval.", "A later agent must actually perform the retrieval."],
        ["Configuration drift", "Deterministic suite with an expected-removal assertion.", "The fresh suite currently detects stale local Resend."],
    ], [1.25, 2.85, 2.40], 8.3)
    p(doc, "The central remediation achievement is not the addition of OpenAI metadata. It is the closed path from canonical identity, through constrained coordination and external-action controls, to a narrow evidence record with clear status boundaries. The repository posture is internally coherent and testable. Current environment conformity remains conditional on reconciling the stale local Resend package.")

    page(doc)
    heading(doc, "Appendix A  Remediated skill inventory")
    inventory = [
        ("aai-cognitive-interface", "Governing state custody and evidence controls"), ("authority-currency-auditor", "Primary-source authority and currency analysis"),
        ("claim-source-auditor", "Atomic claim and source reconciliation"), ("courtlistener-api", "Controlled legal-research API method"),
        ("dify-agent-api", "Dify API with dry-run and execute gates"), ("enformion-go-api", "Credential-free query planning and live disablement"),
        ("exa-firecrawl", "Current web retrieval boundary"), ("forensic-evidentiary-drafting", "Evidentiary drafting method"),
        ("personal-context", "Prior-artifact retrieval with scope control"), ("practitioner-narrative-writer", "Professional narrative drafting control"),
        ("prompt-architecture-engineering", "Prompt design without invented tool claims"), ("record-series-builder", "Record series and state control"),
        ("reddit-automod-yaml", "Static AutoMod and public-copy controls"), ("reddit-owner-ops", "Owner operations with human gates"),
        ("register-mediation", "One-retry mediation workflow"), ("regulatory-complaint-drafting", "Evidence-bound regulatory drafting"),
        ("research-execution-briefs", "Scoped research and source hierarchy"), ("resend-api", "Repository removal and local residue detection"),
        ("speko-mcp", "Origin-constrained API helper"), ("subreddit-rule-packet", "Rule packet generation and validation"),
    ]
    table(doc, ["Skill", "Control role"], inventory, [2.55, 3.95], 8.3)

    page(doc)
    heading(doc, "Appendix B  Reproduction procedure")
    p(doc, "Use the local no-network suite with a temporary report path so the saved repository audit is not overwritten. Inspect prompt mappings, expected exit codes, and removal-state observations. Do not promote its results into claims of installation, activation, provider delivery, or account effect.")
    for text in [
        "python scripts/run_adversarial_skill_suite.py --output <temporary-report>.json",
        "python work/aai-cognitive-interface/scripts/aai_runtime_gate.py package work/aai-cognitive-interface",
        "Inspect the temporary JSON report for every case result and the final status.",
    ]:
        para = doc.add_paragraph(); para.paragraph_format.left_indent = Inches(0.25); para.paragraph_format.first_line_indent = Inches(-0.18); para.paragraph_format.space_after = Pt(3)
        r = para.add_run("•  " + text); run_font(r, size=9.1, name="Courier New" if text.startswith("python") else "Times New Roman")

    page(doc)
    heading(doc, "References")
    refs = [
        "[R1] `.grok/workspace.yaml`, source workspace marker identifying `host: grok`. Inspected 16 September 2026.",
        "[R2] `work/aai-cognitive-interface/SKILL.md` and `references/package-identity.md`, control contract, canonical identity, and status restrictions. Inspected 16 September 2026.",
        "[R3] `work/aai-cognitive-interface/references/host-binding.md`, `runtime-only.md`, and `trusted-controller.md`, host and runtime claim boundaries. Inspected 16 September 2026.",
        "[R4] `.aai-stage-claim-source-auditor/claim-source-auditor/SKILL.md` and `agents/openai.yaml`, source-audit and active-turn coordination contract. Inspected 16 September 2026.",
        "[R5] `.aai-stage-reddit-governed-skills/reddit-owner-ops/SKILL.md` and `agents/openai.yaml`, live-mutation human gates and Codex binding. Inspected 16 September 2026.",
        "[R6] `audits/adversarial-skill-suite-2026-09-16.json` and `scripts/run_adversarial_skill_suite.py`, saved 27-case local evidence and its deterministic method. Inspected 16 September 2026.",
        "[R7] Git commits `874a4bd`, `f7b9b45`, `28ac242`, `8aa09d1`, and `220896f`, remediation lineage and change sets. Inspected 16 September 2026.",
        "[R8] Current temporary execution of `scripts/run_adversarial_skill_suite.py` at 2026-09-16T21:02:02Z. Result: 26 passed and 1 failed because local `resend-api` remains present.",
    ]
    for ref in refs:
        para = p(doc, ref, size=9, after=5); para.paragraph_format.left_indent = Inches(0.23); para.paragraph_format.first_line_indent = Inches(-0.23)

    doc.core_properties.title = "Remediating Grok Origin Skills for OpenAI Compatible Codex Deployment"
    doc.core_properties.subject = "Technical remediation report"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
