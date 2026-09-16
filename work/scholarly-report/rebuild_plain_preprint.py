from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("grok_openai_skill_remediation_technical_report.docx")


def set_font(style, name="Times New Roman", size=10.7, bold=False, italic=False):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)


def add_body(doc, text, indent=0):
    para = doc.add_paragraph(style="Normal")
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.left_indent = Inches(indent)
    para.add_run(text)
    return para


def add_heading(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    para.add_run(text)
    return para


def add_bullet(doc, text):
    para = doc.add_paragraph(style="Normal")
    para.paragraph_format.left_indent = Inches(0.28)
    para.paragraph_format.first_line_indent = Inches(-0.18)
    para.paragraph_format.space_after = Pt(2)
    para.add_run("• " + text)
    return para


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.95)
    sec.right_margin = Inches(0.95)

    normal = doc.styles["Normal"]
    set_font(normal)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)

    title = doc.styles["Title"]
    set_font(title, size=17, bold=True)
    title.paragraph_format.space_after = Pt(8)

    for level, size in ((1, 13), (2, 11.2)):
        style = doc.styles[f"Heading {level}"]
        set_font(style, size=size, bold=True)
        style.paragraph_format.space_before = Pt(15 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.add_run("Remediating Grok Origin Skills for OpenAI Compatible Codex Deployment")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(14)
    run = sub.add_run("Technical remediation account and evidence bounded verification")
    run.italic = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(10.5)

    abstract_label = doc.add_paragraph()
    abstract_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abstract_label.paragraph_format.space_after = Pt(3)
    run = abstract_label.add_run("Abstract")
    run.bold = True
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(11)

    abstract = ("This report explains how a Grok-origin skill workspace was remediated for local Codex use under an OpenAI-compatible control model. The remediation normalized package identity, introduced Codex-facing metadata, constrained coordination to tools actually exposed in the active turn, separated reusable method from mutable world knowledge, and added authorization, source, and evidence boundaries. It also changed external-action helpers so that inspection can occur without credentials or network effect, while live action requires explicit gates or is disabled. Repository evidence records a 27-case local pass. A fresh Codex re-run produced 26 passes and one expected-removal failure because a local `resend-api` package remains installed although the repository remediation removed that route. The report therefore distinguishes a documented repository remediation from a fully reconciled destination environment. It does not claim OpenAI certification, ChatGPT activation, or universal runtime safety.")
    ap = add_body(doc, abstract)
    ap.paragraph_format.left_indent = Inches(0.18)
    ap.paragraph_format.right_indent = Inches(0.18)
    ap.paragraph_format.space_after = Pt(8)

    keywords = doc.add_paragraph()
    keywords.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first = keywords.add_run("Keywords  ")
    first.bold = True
    first.font.name = "Times New Roman"
    second = keywords.add_run("Codex skills, host portability, provenance, authorization gates, runtime validation")
    second.font.name = "Times New Roman"
    keywords.paragraph_format.space_after = Pt(10)

    add_heading(doc, "1  Research question and scope")
    add_body(doc, "The source repository contains a `.grok/workspace.yaml` marker identifying the workspace host as `grok` [R1]. The technical question is not whether a former-host archive can be renamed for a new environment. It is whether the reusable parts of that archive can be converted into an auditable local Codex skill family without importing unsupported host assumptions, historical secrets, stale operational data, or silent authority claims.")
    add_body(doc, "The phrase OpenAI compatible is deliberately narrow. It refers to use of the local Codex skill structure and to controls that align with its visible execution model: `SKILL.md` contracts, `agents/openai.yaml` metadata, active-turn tool discovery, local files, explicit authorization, and deterministic checks. It does not mean that OpenAI has certified, endorsed, installed, or activated the packages.")

    add_heading(doc, "2  Evidence model")
    add_body(doc, "The remediation uses an evidence hierarchy because technical portability often fails when a weak observation is upgraded into a strong conclusion. A commit proves that a repository state was recorded. A package file proves that an instruction or helper exists. A static validator proves the properties it checks for its declared target. A no-network local behavior test proves a particular response to a particular input. None of these proves current provider state, live account authority, external delivery, or activation in another product.")
    add_bullet(doc, "Git history supports remediation lineage, not deployment state.")
    add_bullet(doc, "Package metadata supports intended host binding, not installation or activation.")
    add_bullet(doc, "A static package pass supports structure, not runtime reliability.")
    add_bullet(doc, "A credential-free dry run supports local helper behavior, not provider success or billing outcome.")
    add_bullet(doc, "A blocked operation supports the presence of a gate, not the safety of every untested path.")
    add_body(doc, "This taxonomy is carried through the entire skill family. It prevents a local manifest from becoming evidence of ChatGPT activation and prevents a test artifact from becoming a substitute for current-run verification [R2, R3].")

    add_heading(doc, "3  Remediation architecture")
    add_heading(doc, "3.1  Canonical package identity", 2)
    add_body(doc, "The migration uses canonical skill directories rather than source-host transport wrappers. Each target package has a `SKILL.md` interface contract, a version record, and, where relevant, `agents/openai.yaml` metadata. This establishes an inspectable package identity. The metadata supplies an intended local Codex binding and a default invocation interface. It does not attest that the package has been saved, discovered, or run by a host.")
    add_heading(doc, "3.2  Governor and subordinate methods", 2)
    add_body(doc, "The remediated family makes `aai-cognitive-interface` the governing control layer. Domain modules supply methods only. They do not self-govern, establish a hidden background agent, invent an unseen sibling channel, or turn a draft into an external action. The active agent coordinates only tools exposed in the current turn. The claim-source auditor explicitly treats a sibling result as candidate evidence rather than a competing authority [R4].")
    add_heading(doc, "3.3  Runtime-only boundary", 2)
    add_body(doc, "The packages were converted from possible containers of remembered operational content into runtime overlays. Control rules, status vocabulary, output contracts, and source-location methods can remain in the package. Current legal text, addresses, account state, watch lists, subject information, and dated reference data cannot be treated as current knowledge. They must be retrieved from an authorized current source or reported as unresolved [R3].")
    add_heading(doc, "3.4  Host binding and status honesty", 2)
    add_body(doc, "The host-binding contract names Codex as the local target but makes no activation claim. The package gate can establish a static result only after its defined validator passes. It cannot establish saved, installed, smoke, runtime-verified, or adversarial status without additional evidence. This is a remediation of both language and logic: a package stops at the strongest conclusion its evidence can actually bear.")

    add_heading(doc, "4  External action hardening")
    add_body(doc, "The API and operations modules carried the most important behavioral changes. The remediation introduced a default path that is inspectable, credential-safe, and no-network, followed by an explicit live-action gate. A live flag never replaces the underlying requirement that the user has authorized the exact current action.")
    add_heading(doc, "4.1  Dify Agent API", 2)
    add_body(doc, "The Dify helper was changed to support `--dry-run` and to require `--execute` for a live request. The dry run serializes the intended HTTP method, path, query, and body without loading credentials or contacting Dify. A request without either mode is blocked. This makes local inspection and network execution mechanically distinguishable, and the current no-network suite exercises both outcomes.")
    add_heading(doc, "4.2  EnformionGO API", 2)
    add_body(doc, "The Enformion helper was changed more conservatively. Earlier code could load an access profile and issue an HTTP request. The remediated helper retains a credential-free dry-run plan but blocks every live request until a trusted controller can establish a current, scope-bound authorization receipt. The receipt must bind the action to a subject, lawful purpose, route, minimum fields, and expiry. This prevents a prompt-induced query from becoming a paid or sensitive lookup solely because credentials are locally available.")
    add_heading(doc, "4.3  Speko API", 2)
    add_body(doc, "The Speko helper now requires `--execute`, validates that a path is relative, permits only an explicit method allowlist, pins the destination to HTTPS `api.speko.dev`, and rejects redirects. The ordering is important. Path validation occurs before credential loading, which reduces the opportunity for a malformed argument or redirect chain to move an authorization header to another origin.")
    add_heading(doc, "4.4  Reddit operations and Resend removal", 2)
    add_body(doc, "The Reddit owner-operations module separates drafting and local review from live moderation. Every live write, including a wiki paste, removal, ban, settings change, or modmail act, is a human gate [R5]. The repository remediation also removed the Resend route rather than retaining an email-sending capability whose safe use could not be independently established in the target workflow.")

    add_heading(doc, "5  Scientific and evidentiary controls")
    add_heading(doc, "5.1  Atomic claim assessment", 2)
    add_body(doc, "The claim-source auditor decomposes a source-backed draft into atomic claims and assigns each exactly one result: verified, conflicting, not found in searched sources, or manual review needed. A citation is not enough. The cited material must entail the claim, and a verified or conflicting result must preserve a stable locator. A claim not found in the defined source scope is not treated as false. It remains unconfirmed within that scope.")
    add_heading(doc, "5.2  Failure visibility", 2)
    add_body(doc, "The governing package requires a missing connector, empty result, failed request, or safety interruption to remain visible. A helper must not narrate a completion after an execution boundary stopped it. This is accompanied by status stratification: drafted, static, saved, installed, runtime-smoke, runtime-verified, and adversarial states are distinct claims with distinct evidence requirements. No state may be promoted merely because a related state is present.")
    add_heading(doc, "5.3  Repository lineage", 2)
    add_body(doc, "The implementation history records a sequential remediation rather than a single opaque change. Commit `874a4bd` added the remediated claim-source auditor. Commit `f7b9b45` added governed Reddit skill modules. Commit `28ac242` added governed API and drafting modules. Commit `8aa09d1` hardened helper behavior, added the adversarial suite, and removed the repository Resend package. Commit `220896f` added scholarly critical-thinking prompts. The file sets include contracts, metadata, references, schemas, fixtures, helpers, and test artifacts [R7].")

    add_heading(doc, "6  Verification results")
    add_body(doc, "The saved repository audit, `audits/adversarial-skill-suite-2026-09-16.json`, reports 27 of 27 local cases passing. The suite contains 20 prompt-to-control-marker mappings and seven credential-free dynamic checks. Its declared scope is static controls and local behavior only. It does not call an external API, write to Reddit, or invoke a model provider [R6].")
    add_body(doc, "A fresh execution in the current Codex environment produced 26 passes and one failure. Every retained prompt mapping passed. Dify dry-run and live-gate tests passed. Enformion dry-run and live-disablement tests passed. Speko dry-run and unsafe-origin tests passed. The AutoMod hardening test passed. The remaining failure is the expected-removal check for `resend-api`.")
    add_heading(doc, "6.1  The current Resend discrepancy", 2)
    add_body(doc, "The suite expects `C:\\Users\\jared\\.codex\\skills\\resend-api\\SKILL.md` to be absent because the remediated repository removed that package. The file remains present in the current local skill root. This is a configuration discrepancy between the repository target state and the installed environment. It does not prove a live Resend request, an email delivery, or a failure of unrelated controls. It prevents only the stronger claim that the current environment matches the remediated repository's 27-case target state.")
    add_body(doc, "The appropriate remediation is narrow: under explicit authorization, remove or archive the stale local package and re-run the suite. This report does not perform that deletion because its scope is technical explanation and document production, not mutation of the local skill root.")

    add_heading(doc, "7  Limitations and conclusion")
    add_body(doc, "The remediation provides an auditable local control posture. It does not prove that OpenAI has reviewed or certified the skill family. It does not prove that ChatGPT can access local Codex directories. It does not validate downstream account permissions, provider availability, hidden dependencies, future prompt behavior, or actions outside the tested routes. A current primary-source check and a current authorization remain necessary for every consequential external action.")
    add_body(doc, "Within those limits, the remediation is technically substantial. It translates a source-host workspace into a controlled Codex package family by reducing implied authority, encoding active-turn boundaries, separating current facts from reusable methods, and making dangerous transitions explicit in executable helpers. The repository posture is documented and locally testable. Full present-environment conformity remains conditional on resolving the stale Resend installation.")

    add_heading(doc, "Appendix A  Skill inventory")
    for item in [
        "aai-cognitive-interface: governing state custody and evidence controls.",
        "authority-currency-auditor and claim-source-auditor: primary-source and claim-level verification methods.",
        "courtlistener-api, dify-agent-api, enformion-go-api, exa-firecrawl, and speko-mcp: bounded external-service methods.",
        "forensic-evidentiary-drafting, practitioner-narrative-writer, regulatory-complaint-drafting, and research-execution-briefs: evidence-aware drafting methods.",
        "personal-context, record-series-builder, register-mediation, and prompt-architecture-engineering: retrieval, record, recovery, and prompt-control methods.",
        "reddit-automod-yaml, reddit-owner-ops, and subreddit-rule-packet: local moderation artifact and human-gate methods.",
        "resend-api: removed from the repository target state; current local residue detected by the suite.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Appendix B  Reproduction procedure")
    add_body(doc, "Run the local adversarial suite with a temporary output path so that the saved repository audit is not overwritten. Then inspect each result, including expected exit codes and removal-state checks. The commands below test only the declared local scope and must not be read as activation or provider-side evidence.")
    add_bullet(doc, "python scripts/run_adversarial_skill_suite.py --output <temporary-report>.json")
    add_bullet(doc, "python work/aai-cognitive-interface/scripts/aai_runtime_gate.py package work/aai-cognitive-interface")
    add_bullet(doc, "Review the temporary JSON report before claiming any result beyond its defined local scope.")

    add_heading(doc, "References")
    for ref in [
        "[R1] `.grok/workspace.yaml`, workspace marker identifying `host: grok`. Inspected 16 September 2026.",
        "[R2] `work/aai-cognitive-interface/SKILL.md` and `references/package-identity.md`, control contract and status restrictions. Inspected 16 September 2026.",
        "[R3] `work/aai-cognitive-interface/references/host-binding.md`, `runtime-only.md`, and `trusted-controller.md`, host and runtime claim boundaries. Inspected 16 September 2026.",
        "[R4] `.aai-stage-claim-source-auditor/claim-source-auditor/SKILL.md` and `agents/openai.yaml`, evidence and active-turn coordination contract. Inspected 16 September 2026.",
        "[R5] `.aai-stage-reddit-governed-skills/reddit-owner-ops/SKILL.md` and `agents/openai.yaml`, live-mutation human gates. Inspected 16 September 2026.",
        "[R6] `audits/adversarial-skill-suite-2026-09-16.json` and `scripts/run_adversarial_skill_suite.py`, saved local evidence and deterministic methodology. Inspected 16 September 2026.",
        "[R7] Git commits `874a4bd`, `f7b9b45`, `28ac242`, `8aa09d1`, and `220896f`, remediation lineage. Inspected 16 September 2026.",
        "[R8] Current temporary run at 2026-09-16T21:02:02Z: 26 passed and 1 failed because local `resend-api` remains present.",
    ]:
        para = add_body(doc, ref)
        para.paragraph_format.left_indent = Inches(0.22)
        para.paragraph_format.first_line_indent = Inches(-0.22)
        para.paragraph_format.space_after = Pt(4)
        for run in para.runs:
            run.font.size = Pt(9.2)

    doc.core_properties.title = "Remediating Grok Origin Skills for OpenAI Compatible Codex Deployment"
    doc.core_properties.subject = "Technical remediation report"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
